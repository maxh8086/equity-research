"""Parser for the `Data Sheet` of a Screener.in Excel export (template version 2.1).

Screener fills this sheet with raw numbers; every other sheet in the workbook
is the user's template and is ignored (its formulas are not ours, R1). The
layout is checked strictly: a moved row, a renamed label, a new template
version or a non-number where a number belongs rejects the whole file.

Values are in ₹ crore on the sheet and are stored in absolute rupees.
A blank cell stays absent; it is never read as zero here (the comparison
decides what a blank means, under its own rule version).
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

RULE_VERSION = "screener_export/1"

SHEET = "Data Sheet"
TEMPLATE_VERSION = Decimal("2.1")
CRORE = Decimal(10) ** 7

# Ten periods sit in columns B..K, oldest first.
PERIOD_COLUMNS = tuple(get_column_letter(i) for i in range(2, 12))

INR, SHARES, INR_PER_SHARE = "INR", "shares", "INR_per_share"


@dataclass(frozen=True)
class Section:
    statement: str  # pl | quarter | bs | cf: quarters are three months, the others fiscal years
    date_row: int
    lines: tuple[tuple[str, int, str, str], ...]  # (line, row, column-A label, unit)


LABELS = {1: "COMPANY NAME", 2: "LATEST VERSION", 3: "CURRENT VERSION"}

SECTIONS = (
    Section("pl", 16, (
        ("sales", 17, "Sales", INR),
        ("raw_material", 18, "Raw Material Cost", INR),
        ("inventory_change", 19, "Change in Inventory", INR),
        ("power_fuel", 20, "Power and Fuel", INR),
        ("other_mfr", 21, "Other Mfr. Exp", INR),
        ("employee", 22, "Employee Cost", INR),
        ("selling_admin", 23, "Selling and admin", INR),
        ("other_exp", 24, "Other Expenses", INR),
        ("other_income", 25, "Other Income", INR),
        ("depreciation", 26, "Depreciation", INR),
        ("interest", 27, "Interest", INR),
        ("pbt", 28, "Profit before tax", INR),
        ("tax", 29, "Tax", INR),
        ("net_profit", 30, "Net profit", INR),
        ("dividend", 31, "Dividend Amount", INR),
    )),
    Section("quarter", 41, (
        ("sales", 42, "Sales", INR),
        ("expenses", 43, "Expenses", INR),
        ("other_income", 44, "Other Income", INR),
        ("depreciation", 45, "Depreciation", INR),
        ("interest", 46, "Interest", INR),
        ("pbt", 47, "Profit before tax", INR),
        ("tax", 48, "Tax", INR),
        ("net_profit", 49, "Net profit", INR),
        ("operating_profit", 50, "Operating Profit", INR),
    )),
    Section("bs", 56, (
        ("share_capital", 57, "Equity Share Capital", INR),
        ("reserves", 58, "Reserves", INR),
        ("borrowings", 59, "Borrowings", INR),
        ("other_liabilities", 60, "Other Liabilities", INR),
        ("total_liabilities", 61, "Total", INR),
        ("net_block", 62, "Net Block", INR),
        ("cwip", 63, "Capital Work in Progress", INR),
        ("investments", 64, "Investments", INR),
        ("other_assets", 65, "Other Assets", INR),
        ("total_assets", 66, "Total", INR),
        ("receivables", 67, "Receivables", INR),
        ("inventory", 68, "Inventory", INR),
        ("cash_bank", 69, "Cash & Bank", INR),
        ("share_count", 70, "No. of Equity Shares", SHARES),
        ("new_bonus_shares", 71, "New Bonus Shares", SHARES),
        ("face_value", 72, "Face value", INR_PER_SHARE),
    )),
    Section("cf", 81, (
        ("cfo", 82, "Cash from Operating Activity", INR),
        ("cfi", 83, "Cash from Investing Activity", INR),
        ("cff", 84, "Cash from Financing Activity", INR),
        ("net_cash_flow", 85, "Net Cash Flow", INR),
    )),
)  # fmt: skip


class ExportRejected(Exception):
    """The file is not a Screener Data Sheet this parser understands. Nothing from it is stored."""


@dataclass(frozen=True)
class ScreenerValue:
    statement: str  # pl | quarter | bs | cf
    line: str
    period_end: date
    value: Decimal
    unit: str


@dataclass(frozen=True)
class ParsedExport:
    company_name: str
    template_version: str
    values: tuple[ScreenerValue, ...]

    @property
    def latest_period_end(self) -> date | None:
        return max((v.period_end for v in self.values), default=None)


def _number(cell: object, where: str) -> Decimal | None:
    if cell is None or cell == "":
        return None
    if isinstance(cell, bool) or not isinstance(cell, (int, float)):
        raise ExportRejected(f"{where}: expected a number, got {cell!r}")
    # repr gives the shortest decimal that round-trips the stored double: 1615.51, not 1615.5099999...
    return Decimal(cell) if isinstance(cell, int) else Decimal(repr(cell))


def _dates(ws, row: int) -> dict[str, date]:
    out: dict[str, date] = {}
    for col in PERIOD_COLUMNS:
        cell = ws[f"{col}{row}"].value
        if cell is None:
            continue
        if not isinstance(cell, datetime) or cell.time() != datetime.min.time():
            raise ExportRejected(f"{SHEET}!{col}{row}: expected a report date, got {cell!r}")
        out[col] = cell.date()
    ordered = list(out.values())
    if ordered != sorted(set(ordered)):
        raise ExportRejected(f"{SHEET} row {row}: report dates are not strictly increasing")
    return out


def parse_export(data: bytes) -> ParsedExport:
    try:
        wb = load_workbook(BytesIO(data), data_only=True)
    except (BadZipFile, OSError, KeyError, ValueError) as exc:
        raise ExportRejected(f"not an xlsx workbook: {exc}") from exc
    try:
        if SHEET not in wb.sheetnames:
            raise ExportRejected(f"no {SHEET!r} sheet (sheets: {wb.sheetnames})")
        ws = wb[SHEET]
        for row, label in LABELS.items():
            if ws[f"A{row}"].value != label:
                raise ExportRejected(f"{SHEET}!A{row} is {ws[f'A{row}'].value!r}, expected {label!r}")
        version = _number(ws["B3"].value, f"{SHEET}!B3")
        if version != TEMPLATE_VERSION:
            raise ExportRejected(f"template version {version}, this parser reads {TEMPLATE_VERSION}")
        name = ws["B1"].value
        if not isinstance(name, str) or not name.strip():
            raise ExportRejected(f"{SHEET}!B1 has no company name")

        values: list[ScreenerValue] = []
        for section in SECTIONS:
            if ws[f"A{section.date_row}"].value != "Report Date":
                raise ExportRejected(f"{SHEET}!A{section.date_row} is not 'Report Date'")
            dates = _dates(ws, section.date_row)
            for line, row, label, unit in section.lines:
                if ws[f"A{row}"].value != label:
                    raise ExportRejected(f"{SHEET}!A{row} is {ws[f'A{row}'].value!r}, expected {label!r}")
                for col in PERIOD_COLUMNS:
                    number = _number(ws[f"{col}{row}"].value, f"{SHEET}!{col}{row}")
                    if number is None:
                        continue
                    if col not in dates:
                        raise ExportRejected(f"{SHEET}!{col}{row} has a value but no report date")
                    value = number * CRORE if unit == INR else number
                    values.append(ScreenerValue(section.statement, line, dates[col], value, unit))
    finally:
        wb.close()
    if not values:
        raise ExportRejected(f"{SHEET} has no values")
    return ParsedExport(name.strip(), str(version), tuple(values))
