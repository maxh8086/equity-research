"""The Data Sheet must survive a template build byte for byte.

Screener writes its values into the Data Sheet of whatever workbook it is
handed. openpyxl 3.1 rewrites every string as an inline string, so a cell that
left Screener as a shared string (`t="s"`) comes back as `inlineStr`; Screener
then writes a bare `<v>` into it without clearing the stale `<is>`, Excel
rejects the cell type and discards the sheet, and the export opens blank.
build.py grafts the base workbook's own Data Sheet part back into the saved
zip. These tests pin that down.
"""

from __future__ import annotations

import importlib.util
import posixpath
import re
import shutil
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
BASE_FIXTURE = ROOT / "tests" / "fixtures" / "screener_export" / "vinati_data_sheet_v2_1.xlsx"

MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
DOC_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
SST_PART = "xl/sharedStrings.xml"
HYPERLINK = "https://www.screener.in/excel/"
STYLE_TABLES = ("numFmts", "fonts", "fills", "borders", "cellStyleXfs", "cellXfs", "dxfs")


def _load_build():
    """tools/ is not a package, so load build.py by path."""
    spec = importlib.util.spec_from_file_location(
        "screener_template_build", ROOT / "tools" / "screener_template" / "build.py"
    )
    module = importlib.util.module_from_spec(spec)
    # dataclasses resolves annotations through sys.modules; register before exec.
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


build_mod = _load_build()


# --------------------------------------------------------------------------- #
# The committed fixture is openpyxl's own output, so its strings are already
# inline. Screener ships shared strings, which is the case that broke. Rewrite
# the fixture into that form so the tests exercise the defect.
# --------------------------------------------------------------------------- #
def _shared_string_base(source: Path, out: Path) -> Path:
    with zipfile.ZipFile(source) as zf:
        parts = {i.filename: zf.read(i.filename) for i in zf.infolist()}

    sheet_part = "xl/worksheets/sheet1.xml"
    sheet = parts[sheet_part].decode("utf-8")
    strings: list[str] = []

    def to_shared(m: re.Match[str]) -> str:
        text = m.group("text")
        if text not in strings:
            strings.append(text)
        return f'{m.group("open")} t="s"><v>{strings.index(text)}</v></c>'

    sheet = re.sub(
        r'(?P<open><c\s[^>]*?) t="inlineStr"><is><t>(?P<text>[^<]*)</t></is></c>',
        to_shared,
        sheet,
    )
    assert strings and "inlineStr" not in sheet, "the fixture no longer has inline strings"

    # Screener's Data Sheet carries an external hyperlink, so it has its own rels part.
    sheet = sheet.replace(
        "</worksheet>",
        f'<hyperlinks><hyperlink xmlns:r="{DOC_REL_NS}" ref="E1" display="{HYPERLINK}"'
        ' r:id="rId1"/></hyperlinks></worksheet>',
    )
    parts[sheet_part] = sheet.encode("utf-8")
    parts["xl/worksheets/_rels/sheet1.xml.rels"] = (
        f'<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="{PKG_REL_NS}">'
        f'<Relationship Type="{DOC_REL_NS}/hyperlink" Target="{HYPERLINK}"'
        ' TargetMode="External" Id="rId1" /></Relationships>'
    ).encode("utf-8")

    items = "".join(f"<si><t>{s}</t></si>" for s in strings)
    parts[SST_PART] = (
        f'<?xml version="1.0" encoding="UTF-8"?><sst xmlns="{MAIN_NS}" '
        f'count="{len(strings)}" uniqueCount="{len(strings)}">{items}</sst>'
    ).encode("utf-8")
    parts["xl/_rels/workbook.xml.rels"] = (
        parts["xl/_rels/workbook.xml.rels"]
        .decode("utf-8")
        .replace(
            "</Relationships>",
            f'<Relationship Type="{DOC_REL_NS}/sharedStrings" Target="sharedStrings.xml"'
            ' Id="rIdSst" /></Relationships>',
        )
        .encode("utf-8")
    )
    parts["[Content_Types].xml"] = (
        parts["[Content_Types].xml"]
        .decode("utf-8")
        .replace(
            "</Types>",
            f'<Override PartName="/{SST_PART}" ContentType="application/vnd.'
            'openxmlformats-officedocument.spreadsheetml.sharedStrings+xml" /></Types>',
        )
        .encode("utf-8")
    )

    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in parts.items():
            zf.writestr(name, data)
    return out


def _data_sheet_part(path: Path) -> str:
    with zipfile.ZipFile(path) as zf:
        book = ET.fromstring(zf.read("xl/workbook.xml"))
        sheet = next(s for s in book.iter(f"{{{MAIN_NS}}}sheet") if s.get("name") == "Data Sheet")
        rid = sheet.get(f"{{{DOC_REL_NS}}}id")
        rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
        target = next(r.get("Target", "") for r in rels if r.get("Id") == rid)
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join("xl", target))


def _read(path: Path, part: str) -> bytes:
    with zipfile.ZipFile(path) as zf:
        return zf.read(part)


def _names(path: Path) -> set[str]:
    with zipfile.ZipFile(path) as zf:
        return set(zf.namelist())


def _style_entries(path: Path, tag: str) -> list[bytes]:
    el = ET.fromstring(_read(path, "xl/styles.xml")).find(f"{{{MAIN_NS}}}{tag}")
    return [] if el is None else [ET.tostring(child) for child in el]


@pytest.fixture(scope="module")
def shared_base(tmp_path_factory) -> Path:
    return _shared_string_base(BASE_FIXTURE, tmp_path_factory.mktemp("base") / "base.xlsx")


@pytest.fixture(scope="module")
def built(shared_base: Path, tmp_path_factory) -> Path:
    out = tmp_path_factory.mktemp("built") / "merged.xlsx"
    build_mod.build(shared_base, out)
    return out


def test_data_sheet_part_is_the_base_workbooks_bytes(shared_base: Path, built: Path) -> None:
    assert _read(built, _data_sheet_part(built)) == _read(shared_base, _data_sheet_part(shared_base))


def test_data_sheet_has_no_inline_strings(built: Path) -> None:
    sheet = _read(built, _data_sheet_part(built)).decode("utf-8")
    assert "inlineStr" not in sheet
    assert 't="s"' in sheet


def test_shared_string_table_comes_across(shared_base: Path, built: Path) -> None:
    assert _read(built, SST_PART) == _read(shared_base, SST_PART)

    rels = ET.fromstring(_read(built, "xl/_rels/workbook.xml.rels"))
    hits = [r for r in rels if r.get("Type", "").endswith("/sharedStrings")]
    assert len(hits) == 1 and hits[0].get("Target") == "sharedStrings.xml"

    types = ET.fromstring(_read(built, "[Content_Types].xml"))
    overrides = [
        o for o in types.iter(f"{{{CT_NS}}}Override") if o.get("PartName") == "/" + SST_PART
    ]
    assert len(overrides) == 1
    assert overrides[0].get("ContentType", "").endswith("sharedStrings+xml")


def test_generated_sheets_index_nothing_in_the_shared_table(built: Path) -> None:
    data_sheet = _data_sheet_part(built)
    with zipfile.ZipFile(built) as zf:
        others = [n for n in zf.namelist() if n.startswith("xl/worksheets/sheet") and n != data_sheet]
        assert others, "the build produced no generated sheets"
        for name in others:
            assert 't="s"' not in zf.read(name).decode("utf-8"), name


def test_sheet_relationships_follow_the_data_sheet(built: Path) -> None:
    folder, name = posixpath.split(_data_sheet_part(built))
    part = f"{folder}/_rels/{name}.rels"
    assert part in _names(built)
    assert [r.get("Target") for r in ET.fromstring(_read(built, part))] == [HYPERLINK]


@pytest.mark.parametrize("tag", STYLE_TABLES)
def test_base_style_entries_keep_their_indices(shared_base: Path, built: Path, tag: str) -> None:
    """The grafted sheet's `s=` and `dxfId=` are positions into these tables."""
    base_entries = _style_entries(shared_base, tag)
    assert _style_entries(built, tag)[: len(base_entries)] == base_entries


def test_workbook_still_opens_and_keeps_every_sheet(built: Path) -> None:
    wb = openpyxl.load_workbook(built)
    assert wb.sheetnames[0] == "Summary"
    assert wb.sheetnames[-1] == "Data Sheet"
    assert wb["Data Sheet"]["B1"].value == "VINATI ORGANICS LTD"
    assert wb["Data Sheet"]["A17"].value == "Sales"


def test_an_inline_string_base_still_builds(tmp_path: Path) -> None:
    """A base with no shared-string table must not grow one."""
    out = tmp_path / "merged_inline.xlsx"
    build_mod.build(BASE_FIXTURE, out)
    assert SST_PART not in _names(out)
    assert _read(out, _data_sheet_part(out)) == _read(BASE_FIXTURE, "xl/worksheets/sheet1.xml")


def test_reordered_styles_stop_the_graft(shared_base: Path, built: Path, tmp_path: Path) -> None:
    """If openpyxl ever reorders the style tables the graft must refuse, not corrupt."""
    target = tmp_path / "target.xlsx"
    shutil.copy(built, target)

    styles = ET.fromstring(_read(shared_base, "xl/styles.xml"))
    cell_xfs = styles.find(f"{{{MAIN_NS}}}cellXfs")
    assert cell_xfs is not None and len(cell_xfs) > 1
    cell_xfs[0], cell_xfs[1] = cell_xfs[1], cell_xfs[0]

    reordered = tmp_path / "reordered_base.xlsx"
    with zipfile.ZipFile(shared_base) as src, zipfile.ZipFile(reordered, "w") as dst:
        for info in src.infolist():
            data = (
                ET.tostring(styles)
                if info.filename == "xl/styles.xml"
                else src.read(info.filename)
            )
            dst.writestr(info.filename, data)

    with pytest.raises(SystemExit, match="reordered"):
        build_mod._graft_data_sheet(reordered, target)


def test_a_base_without_a_data_sheet_is_rejected(tmp_path: Path) -> None:
    stripped = tmp_path / "no_data_sheet.xlsx"
    with zipfile.ZipFile(BASE_FIXTURE) as src, zipfile.ZipFile(stripped, "w") as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "xl/workbook.xml":
                data = data.replace(b'name="Data Sheet"', b'name="Something Else"')
            dst.writestr(info.filename, data)

    with pytest.raises(SystemExit, match="Data Sheet"):
        build_mod.build(stripped, tmp_path / "out.xlsx")
