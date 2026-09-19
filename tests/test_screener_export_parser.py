"""Screener Data Sheet parser: contract tests on the recorded fixture and strict rejections. No database."""

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from ingest.screener_export.parser import CRORE, SHEET, ExportRejected, parse_export

FIXTURE = Path(__file__).parent / "fixtures" / "screener_export" / "vinati_data_sheet_v2_1.xlsx"


def _values(parsed) -> dict[tuple[str, str, date], Decimal]:
    return {(v.statement, v.line, v.period_end): v.value for v in parsed.values}


def _edited(**cells) -> bytes:
    wb = load_workbook(FIXTURE)
    for ref, value in cells.items():
        wb[SHEET][ref] = value
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def test_fixture_parses_to_the_recorded_values():
    parsed = parse_export(FIXTURE.read_bytes())
    assert parsed.company_name == "VINATI ORGANICS LTD"
    assert parsed.template_version == "2.1"
    assert len(parsed.values) == 431
    assert parsed.latest_period_end == date(2022, 9, 30)
    values = _values(parsed)
    assert values["pl", "sales", date(2022, 3, 31)] == Decimal("1615.51") * CRORE
    assert values["bs", "share_count", date(2022, 3, 31)] == Decimal(102782050)
    units = {(v.line, v.unit) for v in parsed.values}
    assert ("share_count", "shares") in units and ("face_value", "INR_per_share") in units
    assert all(v.unit == "INR" for v in parsed.values if v.line not in {"share_count", "new_bonus_shares", "face_value"})


def test_blank_cells_stay_absent():
    values = _values(parse_export(_edited(K17=None)))
    assert ("pl", "sales", date(2022, 3, 31)) not in values


@pytest.mark.parametrize(
    "cells, message",
    [
        ({"A17": "Revenue"}, "expected 'Sales'"),
        ({"A1": "NAME"}, "expected 'COMPANY NAME'"),
        ({"B3": 2.2}, "template version"),
        ({"B1": None}, "no company name"),
        ({"K17": "1,615.51"}, "expected a number"),
        ({"K17": True}, "expected a number"),
        ({"K16": None}, "no report date"),
        ({"K16": "Mar 2022"}, "expected a report date"),
        ({"J16": datetime(2023, 3, 31)}, "strictly increasing"),
        ({"A16": "Date"}, "Report Date"),
    ],
)
def test_layout_changes_reject_the_file(cells, message):
    with pytest.raises(ExportRejected, match=message):
        parse_export(_edited(**cells))


def test_not_a_workbook_is_rejected():
    with pytest.raises(ExportRejected, match="not an xlsx"):
        parse_export(b"<html>login</html>")


def test_workbook_without_data_sheet_is_rejected():
    wb = load_workbook(FIXTURE)
    wb[SHEET].title = "Other"
    out = BytesIO()
    wb.save(out)
    with pytest.raises(ExportRejected, match="no 'Data Sheet'"):
        parse_export(out.getvalue())
